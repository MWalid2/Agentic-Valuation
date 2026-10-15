"""Company-to-industry mapping from Damodaran's classification file.

Damodaran assigns every listed company to one of his industry groups, the same
labels as his industry tables. A company is matched by its US exchange ticker;
otherwise the most common industry among US companies with the same SIC code is
proposed. The file is current, not historical, so a past valuation date uses
today's classification; the caller should let the user confirm it.
"""

from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from dcf_agent.data.cache import DataUnavailable

URL = "https://pages.stern.nyu.edu/~adamodar/pc/datasets/indname.xlsx"
SHEET = "By company name"
US_EXCHANGES = {
    "NYSE",
    "NasdaqGS",
    "NasdaqGM",
    "NasdaqCM",
    "NYSEAM",
    "NYSEArca",
    "BATS",
    "Cboe",
    "OTCPK",
    "OTCQX",
    "OTCQB",
    "OTCEM",
    "OTCNO",
}


@dataclass(frozen=True)
class IndustryMatch:
    industry: str
    method: str  # "ticker" or "sic"
    ticker: str
    sic: str | None
    alternatives: tuple[str, ...]
    ref: str
    note: str


def normalize(ticker: str) -> str:
    """SEC uses BRK-B where exchange listings use BRK.B."""
    return ticker.strip().upper().replace("-", ".")


def parse_classification(content: bytes) -> list[dict[str, str]]:
    """US-listed rows: ticker, company, industry, SIC."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True)
        sheet = workbook[SHEET]
    except (KeyError, OSError, ValueError):
        raise DataUnavailable("Unreadable Damodaran industry classification file") from None
    rows: list[dict[str, str]] = []
    header: list[str] | None = None
    for values in sheet.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in values]
        if header is None:
            if "Exchange:Ticker" in cells and "Industry Group" in cells:
                header = cells
            continue
        record = dict(zip(header, cells, strict=False))
        exchange, _, ticker = record.get("Exchange:Ticker", "").partition(":")
        if exchange not in US_EXCHANGES or not ticker or not record.get("Industry Group"):
            continue
        rows.append(
            {
                "ticker": normalize(ticker),
                "company": record.get("Company Name", ""),
                "industry": record["Industry Group"],
                "sic": record.get("SIC Code", "").split(".")[0],
                "country": record.get("Country", ""),
            }
        )
    workbook.close()
    if len(rows) < 1000:
        raise DataUnavailable("Damodaran classification file has too few US companies")
    return rows


def match_industry(
    rows: list[dict[str, Any]],
    ticker: str,
    sic: str | None,
    labels: set[str],
) -> IndustryMatch:
    """Ticker first; SIC majority second; only labels present in the industry tables."""
    wanted = normalize(ticker)
    direct = [r for r in rows if r["ticker"] == wanted and r["industry"] in labels]
    if direct:
        row = direct[0]
        return IndustryMatch(
            industry=row["industry"],
            method="ticker",
            ticker=wanted,
            sic=sic,
            alternatives=(),
            ref=URL,
            note=f"Damodaran classification of {row['company']}",
        )
    if sic:
        counts = Counter(r["industry"] for r in rows if r["sic"] == sic and r["industry"] in labels)
        if counts:
            ranked = [name for name, _ in counts.most_common()]
            share = counts[ranked[0]] / sum(counts.values())
            return IndustryMatch(
                industry=ranked[0],
                method="sic",
                ticker=wanted,
                sic=sic,
                alternatives=tuple(ranked[1:4]),
                ref=URL,
                note=(
                    f"Ticker not classified; most common industry for SIC {sic} "
                    f"({share:.0%} of {sum(counts.values())} US companies); confirm it"
                ),
            )
    raise DataUnavailable(
        f"No Damodaran industry for ticker {ticker} or SIC {sic}; choose one explicitly"
    )
