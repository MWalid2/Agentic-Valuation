"""Formula-driven valuation workbook mirroring the DCF engine.

Every valuation number is a live formula over blue-font inputs on the Inputs
sheet; the engine's results are written separately on Checks so a recalculated
workbook can be compared with the engine to the cent.
"""

from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from dcf_agent.contracts.accounting import NormalizedFinancials
from dcf_agent.contracts.inputs import RawFinancials
from dcf_agent.contracts.modeling import (
    CompsResult,
    ForecastAssumptions,
    ValuationResult,
    WACCResult,
)
from dcf_agent.contracts.provenance import Sourced
from dcf_agent.contracts.workflow import CritiqueReport
from dcf_agent.engine.wacc import UNLEVERED_BETA, WaccInputs

INPUT = Font(color="0000FF")
BOLD = Font(bold=True)
HEADER = PatternFill("solid", fgColor="DDEBF7")
MONEY = "#,##0;[Red]-#,##0"
RATE = "0.00%"
RATIO = "0.000"
PER_SHARE = "#,##0.00"


@dataclass
class Cells:
    """Workbook addresses of the results checked for parity."""

    value_per_share: dict[str, str] = field(default_factory=dict)
    enterprise_value: dict[str, str] = field(default_factory=dict)
    weighted: str = ""
    wacc: str = ""
    # Sheets whose numbers must all be formulas or blue inputs (the no-hard-code check).
    calculation_sheets: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ModelInputs:
    company: str
    ticker: str
    valuation_date: date
    base_revenue: float
    base_label: str
    scenarios: Sequence[ForecastAssumptions]
    results: Sequence[ValuationResult]
    weighted_value: float
    wacc: WACCResult
    wacc_inputs: WaccInputs
    bridge: dict[str, Sourced[float]]
    basic_shares: float
    raw: RawFinancials | None = None
    normalized: NormalizedFinancials | None = None
    notes: Sequence[str] = ()
    # Base-scenario value per share by WACC (rows) x terminal growth (columns).
    sensitivity: (
        tuple[Sequence[float], Sequence[float], Sequence[Sequence[float | None]]] | None
    ) = None
    # Phase 5: the market price, comps, the Critic's report and the changes it made.
    price: float | None = None
    comps: CompsResult | None = None
    critique: CritiqueReport | None = None
    critic_changes: Sequence[str] = ()


def put(
    ws: Worksheet, ref: str, value: object, fmt: str | None = None, font: Font | None = None
) -> None:
    cell = ws[ref]
    cell.value = value
    if fmt:
        cell.number_format = fmt
    if font:
        cell.font = font


def heading(ws: Worksheet, ref: str, text: str) -> None:
    put(ws, ref, text, font=BOLD)
    ws[ref].fill = HEADER


def build_workbook(m: ModelInputs, path: Path) -> Cells:
    wb = Workbook()
    cells = Cells()
    inputs = wb.active
    assert inputs is not None
    inputs.title = "Inputs"
    inputs.column_dimensions["A"].width = 34
    heading(inputs, "A1", f"{m.company} ({m.ticker}) - DCF inputs (blue = input)")
    put(inputs, "A2", "Valuation date")
    put(inputs, "B2", m.valuation_date, "yyyy-mm-dd", INPUT)
    put(inputs, "A3", f"Base revenue ({m.base_label})")
    put(inputs, "B3", m.base_revenue, MONEY, INPUT)
    put(inputs, "A4", "Basic shares outstanding")
    put(inputs, "B4", m.basic_shares, MONEY, INPUT)

    # WACC raw inputs.
    w, wi = m.wacc, m.wacc_inputs
    heading(inputs, "A6", "Cost of capital inputs")
    rows = [
        ("Risk-free rate", wi.risk_free.value, RATE),
        ("Equity risk premium", wi.erp.value, RATE),
        ("Country risk premium", wi.crp.value, RATE),
        (f"Unlevered beta ({wi.industry.name})", wi.industry.values[UNLEVERED_BETA].value, RATIO),
        ("Regression beta (if used)", w.regression_beta.value if w.regression_beta else 0.0, RATIO),
        ("Beta method", w.beta_method, None),
        ("Market value of equity", wi.equity_value.value, MONEY),
        ("Debt incl. capitalized leases", wi.debt_value.value, MONEY),
        ("Marginal tax rate", wi.marginal_tax.value, RATE),
        ("Pre-tax cost of debt", w.pretax_kd.value, RATE),
    ]
    for offset, (label, val, fmt) in enumerate(rows):
        r = 7 + offset
        put(inputs, f"A{r}", label)
        put(inputs, f"B{r}", val, fmt, INPUT)
    # Named positions on Inputs used by formulas.
    rf, erp, crp, beta_u, beta_reg, method, equity, debt, tax, kd = (
        f"Inputs!$B${7 + k}" for k in range(10)
    )

    # Bridge inputs.
    heading(inputs, "A18", "Equity bridge (signed: + adds to equity)")
    bridge_rows = []
    for offset, (name, item) in enumerate(m.bridge.items()):
        r = 19 + offset
        put(inputs, f"A{r}", name.replace("_", " "))
        put(inputs, f"B{r}", item.value, MONEY, INPUT)
        bridge_rows.append(r)
    bridge_end = 19 + max(len(m.bridge), 1) - 1
    bridge_sum = f"SUM(Inputs!$B$19:$B${bridge_end})" if m.bridge else "0"

    # Scenario drivers.
    driver_start = bridge_end + 3
    scenario_rows: dict[str, dict[str, int]] = {}
    r = driver_start
    for a in m.scenarios:
        heading(inputs, f"A{r}", f"Scenario: {a.scenario}")
        put(inputs, f"B{r}", "Year end")
        for k, end in enumerate(a.year_ends):
            put(inputs, f"{get_column_letter(3 + k)}{r}", end, "yyyy-mm-dd", INPUT)
        positions = {"year_end": r}
        for label, key, series in (
            ("Revenue growth", "growth", a.revenue_growth),
            ("Operating margin", "margin", a.operating_margin),
            ("Sales to capital", "s2c", a.sales_to_capital),
            ("Tax rate", "tax", a.tax_rate),
        ):
            r += 1
            put(inputs, f"A{r}", label)
            for k, item in enumerate(series):
                fmt = RATIO if key == "s2c" else RATE
                put(inputs, f"{get_column_letter(3 + k)}{r}", item.value, fmt, INPUT)
            positions[key] = r
        for label, key, val, fmt in (
            ("Terminal growth", "g", a.terminal_growth.value, RATE),
            ("Terminal ROIC", "roic", a.terminal_roic.value, RATE),
            ("Marginal tax (terminal)", "mtax", a.marginal_tax_rate.value, RATE),
            ("Probability", "p", a.probability, RATE),
        ):
            r += 1
            put(inputs, f"A{r}", label)
            put(inputs, f"B{r}", val, fmt, INPUT)
            positions[key] = r
        scenario_rows[a.scenario] = positions
        r += 2
    for col in range(2, 3 + max(len(a.year_ends) for a in m.scenarios)):
        inputs.column_dimensions[get_column_letter(col)].width = 14

    # WACC build-up.
    ws = wb.create_sheet("WACC")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    heading(ws, "A1", "Cost of capital (formulas)")
    formulas = [
        ("Debt / equity (market)", f"={debt}/{equity}", RATIO),
        ("Relevered bottom-up beta", f"={beta_u}*(1+(1-{tax})*B2)", RATIO),
        ("Beta used", f'=IF({method}="regression",{beta_reg},B3)', RATIO),
        ("Cost of equity", f"={rf}+B4*{erp}+{crp}", RATE),
        ("Pre-tax cost of debt", f"={kd}", RATE),
        ("After-tax cost of debt", f"=B6*(1-{tax})", RATE),
        ("Weight of equity", f"={equity}/({equity}+{debt})", RATE),
        ("Weight of debt", "=1-B8", RATE),
        ("WACC", "=B8*B5+B9*B7", RATE),
        ("Terminal beta", "=IF(B4>1,1,B4)", RATIO),
        ("Terminal WACC", f"=B8*({rf}+B11*{erp}+{crp})+B9*B7", RATE),
    ]
    for offset, (label, formula, fmt) in enumerate(formulas):
        row = 2 + offset
        put(ws, f"A{row}", label)
        put(ws, f"B{row}", formula, fmt)
    wacc_cell, terminal_wacc_cell = "WACC!$B$10", "WACC!$B$12"
    cells.wacc = "WACC!B10"
    for offset, note in enumerate(m.wacc.method_notes):
        put(ws, f"A{15 + offset}", note)

    # One DCF sheet per scenario.
    for a in m.scenarios:
        pos = scenario_rows[a.scenario]
        ws = wb.create_sheet(f"DCF {a.scenario}")
        ws.column_dimensions["A"].width = 34
        n = len(a.year_ends)
        cols = [get_column_letter(3 + k) for k in range(n)]
        last = cols[-1]
        heading(ws, "A1", f"FCFF valuation - {a.scenario} scenario")
        labels = [
            "Year end",
            "Revenue",
            "Operating margin",
            "EBIT",
            "Tax rate",
            "NOPAT",
            "Reinvestment",
            "FCFF (full year)",
            "Stub-adjusted FCFF",
            "Discount time",
            "Discount factor",
            "PV of FCFF",
        ]
        for offset, label in enumerate(labels):
            put(ws, f"A{2 + offset}", label)
        put(ws, "B3", "=Inputs!$B$3", MONEY)  # base revenue in column B
        put(ws, "A16", "Stub fraction")
        put(ws, "B16", f"=(Inputs!$C${pos['year_end']}-Inputs!$B$2)/365", RATIO)
        for k, c in enumerate(cols):
            prev = "B" if k == 0 else cols[k - 1]
            src = get_column_letter(3 + k)
            put(ws, f"{c}2", f"=Inputs!{src}${pos['year_end']}", "yyyy-mm-dd")
            put(ws, f"{c}3", f"={prev}3*(1+Inputs!{src}${pos['growth']})", MONEY)
            put(ws, f"{c}4", f"=Inputs!{src}${pos['margin']}", RATE)
            put(ws, f"{c}5", f"={c}3*{c}4", MONEY)
            put(ws, f"{c}6", f"=Inputs!{src}${pos['tax']}", RATE)
            put(ws, f"{c}7", f"={c}5*(1-{c}6)", MONEY)
            put(ws, f"{c}8", f"=({c}3-{prev}3)/Inputs!{src}${pos['s2c']}", MONEY)
            put(ws, f"{c}9", f"={c}7-{c}8", MONEY)
            put(ws, f"{c}10", f"={c}9*$B$16" if k == 0 else f"={c}9", MONEY)
            put(ws, f"{c}11", "=$B$16/2" if k == 0 else f"=$B$16+{k - 1}+0.5", RATIO)
            put(ws, f"{c}12", f"=(1+{wacc_cell})^(-{c}11)", "0.000000")
            put(ws, f"{c}13", f"={c}10*{c}12", MONEY)
        g, roic, mtax = (f"Inputs!$B${pos[k]}" for k in ("g", "roic", "mtax"))
        block = [
            ("Terminal NOPAT", f"={last}3*(1+{g})*{last}4*(1-{mtax})", MONEY),
            ("Terminal FCFF", f"=B18*(1-{g}/{roic})", MONEY),
            ("Terminal value", f"=B19/({terminal_wacc_cell}-{g})", MONEY),
            ("Terminal discount time", f"=MAX($B$16+{n}-1-0.5,{last}11)", RATIO),
            ("PV of terminal value", f"=B20*(1+{wacc_cell})^(-B21)", MONEY),
            ("PV of explicit FCFF", f"=SUM(C13:{last}13)", MONEY),
            ("Enterprise value", "=B23+B22", MONEY),
            ("Equity bridge (signed)", f"={bridge_sum}", MONEY),
            ("Equity value", "=B24+B25+D26", MONEY),
            ("Shares", "=D27", MONEY),
            ("Value per share", "=B26/B27", PER_SHARE),
            ("Terminal share of EV", "=IF(B24<>0,B22/B24,0)", RATE),
        ]
        for offset, (label, formula, fmt) in enumerate(block):
            row = 18 + offset
            put(ws, f"A{row}", label, font=BOLD if label == "Value per share" else None)
            put(ws, f"B{row}", formula, fmt)
        # Dilution is a fixed point (options by the treasury stock method at the value per
        # share, convertibles if converted): the engine's result is an input here.
        result = next((r for r in m.results if r.scenario == a.scenario), None)
        engine_bridge = sum(v.value for v in m.bridge.values())
        converted = (
            result.equity_value - (result.enterprise_value + engine_bridge) if result else 0.0
        )
        diluted = result.diluted_shares if result else m.basic_shares
        put(ws, "C26", "Convertible principal converted (engine)")
        put(ws, "D26", converted, MONEY, INPUT)
        put(ws, "C27", "Diluted shares (engine: units, options by TSM, convertibles)")
        put(ws, "D27", diluted, MONEY, INPUT)
        for c in ["B", *cols]:
            ws.column_dimensions[c].width = 16
        cells.value_per_share[a.scenario] = f"'DCF {a.scenario}'!B28"
        cells.enterprise_value[a.scenario] = f"'DCF {a.scenario}'!B24"

    # Scenario weighting.
    ws = wb.create_sheet("Scenarios")
    ws.column_dimensions["A"].width = 24
    heading(ws, "A1", "Scenario")
    put(ws, "B1", "Probability", font=BOLD)
    put(ws, "C1", "Value per share", font=BOLD)
    for k, a in enumerate(m.scenarios):
        row = 2 + k
        put(ws, f"A{row}", a.scenario)
        put(ws, f"B{row}", f"=Inputs!$B${scenario_rows[a.scenario]['p']}", RATE)
        put(ws, f"C{row}", f"={cells.value_per_share[a.scenario]}", PER_SHARE)
    last_row = 1 + len(m.scenarios)
    put(ws, f"A{last_row + 1}", "Probability-weighted value", font=BOLD)
    put(ws, f"C{last_row + 1}", f"=SUMPRODUCT(B2:B{last_row},C2:C{last_row})", PER_SHARE)
    cells.weighted = f"Scenarios!C{last_row + 1}"

    # History and normalized metrics as sourced data.
    if m.raw is not None:
        ws = wb.create_sheet("Historical")
        periods = [*m.raw.annual, *((m.raw.ttm,) if m.raw.ttm else ())]
        heading(ws, "A1", "Concept")
        ws.column_dimensions["A"].width = 30
        for k, p in enumerate(periods):
            put(ws, f"{get_column_letter(2 + k)}1", f"{p.kind} {p.period_end}", font=BOLD)
        concepts = sorted({c for p in periods for c in p.values})
        for row, concept in enumerate(concepts, 2):
            put(ws, f"A{row}", concept)
            for k, p in enumerate(periods):
                item = p.values.get(concept)
                status = p.statuses.get(concept)
                put(
                    ws,
                    f"{get_column_letter(2 + k)}{row}",
                    item.value if item else (status.status if status else None),
                    MONEY if item else None,
                )
    if m.normalized is not None:
        ws = wb.create_sheet("Normalized")
        periods_n = [*m.normalized.annual, *((m.normalized.ttm,) if m.normalized.ttm else ())]
        heading(ws, "A1", "Metric")
        ws.column_dimensions["A"].width = 26
        metrics = (
            "revenue",
            "clean_ebit",
            "nopat",
            "invested_capital",
            "effective_tax_rate",
            "roic",
            "sales_to_capital",
            "reinvestment",
            "reinvestment_rate",
        )
        for k, p in enumerate(periods_n):
            put(ws, f"{get_column_letter(2 + k)}1", f"{p.kind} {p.period_end}", font=BOLD)
        for row, metric in enumerate(metrics, 2):
            put(ws, f"A{row}", metric)
            for k, p in enumerate(periods_n):
                item = getattr(p, metric)
                fmt = (
                    RATE if metric in ("effective_tax_rate", "roic", "reinvestment_rate") else MONEY
                )
                put(ws, f"{get_column_letter(2 + k)}{row}", item.value if item else None, fmt)

    if m.sensitivity is not None:
        waccs, growths, grid = m.sensitivity
        ws = wb.create_sheet("Sensitivity")
        ws.column_dimensions["A"].width = 18
        heading(ws, "A1", "Base value per share: WACC (rows) x terminal growth (columns)")
        put(ws, "A2", "Engine-computed values; blank cells violate a guardrail (e.g. WACC <= g)")
        for k, growth in enumerate(growths):
            put(ws, f"{get_column_letter(2 + k)}3", growth, RATE, BOLD)
        for row, (rate, values) in enumerate(zip(waccs, grid, strict=True), 4):
            put(ws, f"A{row}", rate, RATE, BOLD)
            for k, v in enumerate(values):
                put(ws, f"{get_column_letter(2 + k)}{row}", v, PER_SHARE)

    # Engine results for parity.
    ws = wb.create_sheet("Checks")
    ws.column_dimensions["A"].width = 32
    heading(ws, "A1", "Check")
    for letter, title in zip("BCD", ("Engine", "Workbook", "Difference"), strict=True):
        put(ws, f"{letter}1", title, font=BOLD)
    row = 2
    for result in m.results:
        for label, engine, ref, fmt in (
            (
                f"{result.scenario} value per share",
                result.value_per_share,
                cells.value_per_share[result.scenario],
                PER_SHARE,
            ),
            (
                f"{result.scenario} enterprise value",
                result.enterprise_value,
                cells.enterprise_value[result.scenario],
                MONEY,
            ),
        ):
            put(ws, f"A{row}", label)
            put(ws, f"B{row}", engine, fmt)
            put(ws, f"C{row}", f"={ref}", fmt)
            put(ws, f"D{row}", f"=C{row}-B{row}", "0.0000")
            row += 1
    put(ws, f"A{row}", "Probability-weighted value per share")
    put(ws, f"B{row}", m.weighted_value, PER_SHARE)
    put(ws, f"C{row}", f"={cells.weighted}", PER_SHARE)
    put(ws, f"D{row}", f"=C{row}-B{row}", "0.0000")
    put(ws, f"A{row + 1}", "WACC")
    put(ws, f"B{row + 1}", m.wacc.wacc.value, RATE)
    put(ws, f"C{row + 1}", f"={cells.wacc}", RATE)
    put(ws, f"D{row + 1}", f"=C{row + 1}-B{row + 1}", "0.000000")

    if m.critique is not None:
        row += 3
        heading(ws, f"A{row}", "Critic")
        put(
            ws,
            f"B{row}",
            f"score {m.critique.score:.1f} of 5; "
            f"{'passed' if m.critique.passed else 'not passed'}; "
            f"{m.critique.revision_count} loops",
        )
        for column, title in zip("ABCD", ("Severity", "Issue", "Owner", "Fix"), strict=True):
            put(ws, f"{column}{row + 1}", title, font=BOLD)
        for k, issue in enumerate(m.critique.issues, row + 2):
            for column, text in zip(
                "ABCD",
                (issue.severity, issue.description, issue.route_back_to, issue.suggested_fix),
                strict=True,
            ):
                put(ws, f"{column}{k}", text)
        row += 2 + len(m.critique.issues)
        for k, change in enumerate(m.critic_changes, row + 1):
            put(ws, f"A{k}", "Critic change")
            put(ws, f"B{k}", change)

    comps_range: tuple[str, str] | None = None
    if m.comps is not None:
        comps_range = comps_sheet(wb.create_sheet("Comps"), m.comps)
    cover_sheet(wb, m, cells, comps_range)
    cells.calculation_sheets = [
        "Cover",
        "WACC",
        "Scenarios",
        *(f"DCF {a.scenario}" for a in m.scenarios),
    ]

    # Sources and notes.
    ws = wb.create_sheet("Sources")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    heading(ws, "A1", "Item")
    put(ws, "B1", "Source / note", font=BOLD)
    entries: list[tuple[str, str]] = [
        ("Risk-free rate", wi.risk_free.prov.ref),
        ("Equity risk premium", wi.erp.prov.ref),
        ("Industry beta", wi.industry.values[UNLEVERED_BETA].prov.ref),
        ("Market value of equity", wi.equity_value.prov.ref),
        *((f"Bridge: {k}", v.prov.ref[:300]) for k, v in m.bridge.items()),
        *(("Note", n) for n in m.notes),
        ("Disclaimer", "Educational model; not investment advice. Review all inputs."),
    ]
    for row, (label, text) in enumerate(entries, 2):
        put(ws, f"A{row}", label)
        put(ws, f"B{row}", text)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return cells


def comps_sheet(ws: Worksheet, comps: CompsResult) -> tuple[str, str] | None:
    """Peers' multiples and the engine's quartiles (after outliers); the implied range."""
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["F"].width = 90
    heading(ws, "A1", "Comparable companies (TTM multiples; engine data, blue = input)")
    metrics = ("ev_ebitda", "ev_sales", "ev_ebit", "price_earnings")
    for column, title in zip(
        "ABCDEF", ("Peer", "EV/EBITDA", "EV/Sales", "EV/EBIT", "P/E", "Why kept"), strict=True
    ):
        put(ws, f"{column}2", title, font=BOLD)
    for k, peer in enumerate(comps.peers, 3):
        put(ws, f"A{k}", peer.ticker)
        for column, metric in zip("BCDE", metrics, strict=True):
            item = getattr(peer, metric)
            if item is not None:
                put(ws, f"{column}{k}", item.value, "0.0x", INPUT)
        put(ws, f"F{k}", peer.selection_rationale[:300])
    row = 4 + len(comps.peers)
    heading(ws, f"A{row}", "Quartiles after outliers (engine)")
    for column, title in zip("BCDEF", ("Metric", "Peers", "Q1", "Median", "Q3"), strict=True):
        put(ws, f"{column}{row}", title, font=BOLD)
    for k, summary in enumerate(comps.summaries, row + 1):
        put(ws, f"B{k}", summary.metric)
        put(ws, f"C{k}", summary.sample_size, None, INPUT)
        put(ws, f"D{k}", summary.first_quartile, "0.0x", INPUT)
        put(ws, f"E{k}", summary.median, "0.0x", INPUT)
        put(ws, f"F{k}", summary.third_quartile, "0.0x", INPUT)
    row += 2 + len(comps.summaries)
    for k, warning in enumerate(comps.warnings, row + 3):
        put(ws, f"A{k}", "Warning")
        put(ws, f"B{k}", warning)
    if comps.implied_value_range is None:
        return None
    put(ws, f"A{row}", "Implied value per share (low, high)")
    put(ws, f"B{row}", comps.implied_value_range[0], PER_SHARE, INPUT)
    put(ws, f"C{row}", comps.implied_value_range[1], PER_SHARE, INPUT)
    return f"Comps!$B${row}", f"Comps!$C${row}"


def cover_sheet(
    wb: Workbook, m: ModelInputs, cells: Cells, comps_range: tuple[str, str] | None
) -> None:
    """The first sheet: the answer, every number a formula over the model or a blue input."""
    ws = wb.create_sheet("Cover", 0)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    heading(ws, "A1", f"{m.company} ({m.ticker}): DCF valuation")
    price = f"B{4 + len(m.scenarios)}"
    gap = f'=IF(ISNUMBER({price}),B3/{price}-1,"n/a")'
    rows: list[tuple[str, object, str | None, Font | None]] = [
        ("Valuation date", "=Inputs!B2", "yyyy-mm-dd", None),
        ("Probability-weighted value per share", f"={cells.weighted}", PER_SHARE, BOLD),
        *(
            (
                f"{a.scenario.capitalize()} value per share",
                f"={cells.value_per_share[a.scenario]}",
                PER_SHARE,
                None,
            )
            for a in m.scenarios
        ),
        ("Market price", m.price if m.price is not None else "n/a", PER_SHARE, INPUT),
        ("Value against price", gap, "+0%;-0%", None),
        ("WACC", f"={cells.wacc}", RATE, None),
    ]
    if comps_range is not None:
        rows += [
            ("Comps implied value, low", f"={comps_range[0]}", PER_SHARE, None),
            ("Comps implied value, high", f"={comps_range[1]}", PER_SHARE, None),
        ]
    for k, (label, value, fmt, font) in enumerate(rows, 2):
        put(ws, f"A{k}", label)
        put(ws, f"B{k}", value, fmt, font)
    row = 3 + len(rows)
    if m.critique is not None:
        put(ws, f"A{row}", "Critic score (of 5; issues on Checks)")
        put(ws, f"B{row}", m.critique.score, "0.0", INPUT)
        put(ws, f"A{row + 1}", "Critic review")
        put(ws, f"B{row + 1}", "passed" if m.critique.passed else "not passed")
        row += 3
    put(ws, f"A{row}", "Educational model; not investment advice. Sources on the Sources sheet.")
