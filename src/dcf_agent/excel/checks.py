"""Structural checks of a written workbook: formulas where calculations are.

- formula count: every scenario's DCF sheet has its full set of formulas;
- key cells (value per share, enterprise value, weighted value, WACC) are formulas;
- no hard-coded calculation cell: on the calculation sheets every number is a formula
  or a blue-font input (an input is a value the model takes, sourced on Sources).
"""

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from dcf_agent.excel.workbook import INPUT, Cells

# Per DCF sheet: 12 rows a year, base revenue, stub, 12 terminal and bridge rows.
PER_YEAR, PER_SHEET = 12, 14


@dataclass
class Structure:
    formulas: int
    expected: int
    missing_key: list[str] = field(default_factory=list)
    hardcoded: list[str] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return self.formulas >= self.expected and not self.missing_key and not self.hardcoded

    def text(self) -> str:
        return (
            f"formulas {self.formulas} (at least {self.expected}); key cells "
            f"{'all formulas' if not self.missing_key else self.missing_key}; hard-coded "
            f"{'none' if not self.hardcoded else self.hardcoded[:5]}"
        )


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def check_structure(path: Path, cells: Cells, scenarios: int, years: int) -> Structure:
    workbook = load_workbook(path)
    try:
        formulas = sum(
            is_formula(c.value) for ws in workbook.worksheets for row in ws.iter_rows() for c in row
        )
        keys = [
            *cells.value_per_share.values(),
            *cells.enterprise_value.values(),
            cells.weighted,
            cells.wacc,
        ]
        missing = []
        for ref in keys:
            sheet, cell = ref.rsplit("!", 1)
            if not is_formula(workbook[sheet.strip("'")][cell.replace("$", "")].value):
                missing.append(ref)
        hardcoded = []
        for name in cells.calculation_sheets:
            for row in workbook[name].iter_rows():
                for c in row:
                    number = isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
                    if number and (c.font is None or c.font.color != INPUT.color):
                        hardcoded.append(f"{name}!{c.coordinate}")
    finally:
        workbook.close()
    expected = scenarios * (PER_YEAR * years + PER_SHEET)
    return Structure(formulas, expected, missing, hardcoded)
