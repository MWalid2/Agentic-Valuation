"""Reporter: summary, football field, memo, workbook Cover/Comps/Checks and its structure."""

import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_critic import clean
from test_valuation_graph import BUNDLE

from dcf_agent.agents.critic import review
from dcf_agent.contracts.modeling import CompsResult, MultipleSummary, PeerMultiples
from dcf_agent.contracts.provenance import Provenance, Sourced
from dcf_agent.engine.valuation import write_workbook
from dcf_agent.excel.checks import check_structure
from dcf_agent.graph.workflow import MARGINAL_TAX
from dcf_agent.report.chart import football_field
from dcf_agent.report.export import report_files
from dcf_agent.report.summary import Range, build_summary, headline

pytestmark = pytest.mark.skipif(not BUNDLE.exists(), reason="reference bundle not collected")


def multiple(value: float, as_of) -> Sourced[float]:
    return Sourced[float](
        value=value,
        prov=Provenance(
            source="derived",
            ref="dcf_agent.agents.comps",
            as_of=as_of,
            confidence="medium",
            note="EV / EBITDA, TTM",
        ),
    )


def with_comps_and_review():
    state = clean()
    as_of = state["valuation"].as_of
    peers = tuple(
        PeerMultiples(
            ticker=t,
            cik=f"{k + 1:010d}",
            as_of=as_of,
            selection_rationale="Beverage (Soft); test peer",
            ev_ebitda=multiple(v, as_of),
        )
        for k, (t, v) in enumerate((("PEP", 14.0), ("KDP", 12.0), ("MNST", 25.0)))
    )
    state["comps"] = CompsResult(
        as_of=as_of,
        peers=peers,
        summaries=(
            MultipleSummary(
                metric="ev_ebitda", sample_size=3, first_quartile=13, median=14, third_quartile=19.5
            ),
        ),
        implied_value_range=(40.0, 60.0),
        divergence_vs_dcf=-0.1,
        calculation_ref="dcf_agent.agents.comps",
        warnings=("only 3 peers with multiples (fewer than 5)",),
    )
    state["critique"] = review(state)
    return state


def test_the_summary_ranks_drivers_and_takes_numbers_from_the_engine():
    state = with_comps_and_review()
    summary = build_summary(state, ["Changes in consumer preferences could hurt us. More text."])
    valuation = state["valuation"]
    assert summary.weighted == valuation.weighted and summary.low == valuation.low
    assert len(summary.drivers) == 3
    assert summary.drivers[0].width >= summary.drivers[1].width >= summary.drivers[2].width
    base = valuation.scenario("base").value_per_share
    for swing in summary.drivers:  # a driver moved both ways brackets the base value
        assert min(swing.low, swing.high) <= base + 0.01 <= max(swing.low, swing.high) + 0.02
    assert any("10-K risk factors" in r for r in summary.risks)
    assert [r.label.split(" ")[0] for r in summary.ranges][:2] == ["DCF", "Comps"]
    assert headline("First sentence. Second one.") == "First sentence."


def test_the_football_field_is_valid_svg_with_every_range_and_the_price():
    svg = football_field(
        [Range("DCF", 30, 60, 45), Range("Comps", 40, 55)], 50.0, "KO: value per share"
    )
    root = ET.fromstring(svg)
    assert root.tag.endswith("svg")
    assert "$30.00" in svg and "$60.00" in svg and "price $50.00" in svg
    assert football_field([], None, "x") == ""


def test_the_report_writes_memo_chart_and_summary(tmp_path: Path):
    state = with_comps_and_review()
    summary, files = report_files(state, tmp_path, unreviewed=("Gate 1 auto-approved",))
    md = files["memo_md"].read_text(encoding="utf-8")
    page = files["memo_html"].read_text(encoding="utf-8")
    for title in ("## Summary", "## Scenarios", "## What drives the value", "## Cost of capital"):
        assert title in md
    assert "## Comparable companies" in md and "## Unreviewed items (autopilot)" in md
    assert f"${summary.weighted:,.2f}" in md and "![Football field](football-field.svg)" in md
    assert "<svg" in page and "<table>" in page and "<script" not in page
    ET.fromstring(files["chart"].read_text(encoding="utf-8"))
    assert '"weighted"' in files["summary"].read_text(encoding="utf-8")
    words = len(md.split())
    assert 500 < words < 3000  # two to four pages


def test_the_workbook_has_cover_comps_and_critic_and_no_hardcoded_calculation(tmp_path: Path):
    state = with_comps_and_review()
    scenarios = state["forecast"]
    path, cells = write_workbook(
        state["data"],
        scenarios,
        state["wacc"],
        state["valuation"],
        state["normalized"],
        MARGINAL_TAX,
        tmp_path / "m.xlsx",
        comps=state["comps"],
        critique=state["critique"],
    )
    book = load_workbook(path)
    assert book.sheetnames[0] == "Cover" and "Comps" in book.sheetnames
    assert book["Cover"]["B3"].value == f"={cells.weighted}"
    checks = [c.value for row in book["Checks"].iter_rows() for c in row]
    assert "Critic" in checks
    book.close()
    structure = check_structure(path, cells, 3, len(scenarios.scenarios[0].year_ends))
    assert structure.passed, structure.text()
    # A number typed over a formula is caught.
    book = load_workbook(path)
    book["DCF base"]["C5"] = 123.0
    book.save(path)
    broken = check_structure(path, cells, 3, len(scenarios.scenarios[0].year_ends))
    assert not broken.passed and "DCF base!C5" in broken.hardcoded
