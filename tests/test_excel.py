"""Workbook formulas mirror the engine; LibreOffice parity when it is installed."""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_dcf_engine import VALUED, assumptions
from test_wacc import inputs

from dcf_agent.contracts.external import BenchmarkRow
from dcf_agent.contracts.provenance import Provenance, Sourced
from dcf_agent.engine.dcf import Shares, run_dcf, weighted_value
from dcf_agent.engine.wacc import compute_wacc
from dcf_agent.excel.parity import find_soffice, recalculate, verify
from dcf_agent.excel.workbook import ModelInputs, build_workbook


def dated(value: float) -> Sourced[float]:
    return Sourced[float](
        value=value,
        prov=Provenance(source="damodaran", ref="test", as_of=VALUED, confidence="high"),
    )


def model() -> tuple[ModelInputs, list]:
    # Evidence must not follow the valuation date (the contracts reject look-ahead).
    wacc_inputs = inputs(
        as_of=VALUED,
        risk_free=dated(0.04),
        erp=dated(0.05),
        crp=dated(0.0),
        marginal_tax=dated(0.25),
        industry=BenchmarkRow(
            name="Test industry",
            values={"Unlevered beta corrected for cash": dated(0.8)},
            units={"Unlevered beta corrected for cash": "ratio"},
        ),
        equity_value=dated(800.0),
        debt_value=dated(200.0),
    )
    wacc = compute_wacc(wacc_inputs)
    bridge = {
        "cash": Sourced[float](
            value=50.0,
            prov=Provenance(
                source="sec_xbrl",
                ref="cash",
                as_of=VALUED,
                confidence="high",
                filed=date(2025, 2, 1),
            ),
        ),
        "debt": Sourced[float](
            value=-200.0,
            prov=Provenance(
                source="sec_xbrl",
                ref="debt",
                as_of=VALUED,
                confidence="high",
                filed=date(2025, 2, 1),
            ),
        ),
    }
    forecasts = [
        assumptions("bear", 0.25, 0.0),
        assumptions("base", 0.5, 0.10),
        assumptions("bull", 0.25, 0.2),
    ]
    results = [
        run_dcf(
            f,
            wacc,
            base_revenue=100.0,
            bridge=bridge,
            shares=Shares(basic=10.0),
            valuation_date=VALUED,
        )
        for f in forecasts
    ]
    weighted = weighted_value(results, {f.scenario: f.probability for f in forecasts})
    return (
        ModelInputs(
            company="Example",
            ticker="EXM",
            valuation_date=VALUED,
            base_revenue=100.0,
            base_label="annual",
            scenarios=forecasts,
            results=results,
            weighted_value=weighted,
            wacc=wacc,
            wacc_inputs=wacc_inputs,
            bridge=bridge,
            basic_shares=10.0,
        ),
        results,
    )


def test_workbook_has_live_formulas_and_blue_inputs(tmp_path: Path):
    m, _ = model()
    path = tmp_path / "model.xlsx"
    cells = build_workbook(m, path)
    wb = load_workbook(path)
    assert {"Inputs", "WACC", "DCF base", "Scenarios", "Checks", "Sources"} <= set(wb.sheetnames)
    dcf = wb["DCF base"]
    assert str(dcf["B28"].value).startswith("=")  # value per share is a formula
    assert str(dcf["C3"].value).startswith("=")  # revenue projection is a formula
    assert wb["Inputs"]["B3"].font.color.rgb.endswith("0000FF")  # inputs in blue
    assert cells.value_per_share["base"] == "'DCF base'!B28"


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_recalculated_workbook_equals_the_engine_to_the_cent(tmp_path: Path):
    m, results = model()
    path = tmp_path / "model.xlsx"
    cells = build_workbook(m, path)
    recalculated = recalculate(path, tmp_path / "recalculated")
    checks = verify(
        recalculated,
        cells,
        value_per_share={r.scenario: r.value_per_share for r in results},
        enterprise_value={r.scenario: r.enterprise_value for r in results},
        weighted=m.weighted_value,
        wacc=m.wacc.wacc.value,
    )
    assert all(c.passed for c in checks), [(c.label, c.engine, c.workbook) for c in checks]


def test_cli_refuses_a_bundle_that_readiness_blocks(tmp_path: Path, capsys):
    from dcf_agent.cli import main

    bundle = Path("outputs/coca-cola-data.json")
    if not bundle.exists():
        pytest.skip("Reference bundle not built")
    data = bundle.read_text(encoding="utf-8")
    import json

    payload = json.loads(data)
    payload["readiness"]["wacc"] = {
        "status": "blocked",
        "missing": {"ttm": ["interest_expense (unresolved)"]},
        "review_notes": [],
    }
    blocked = tmp_path / "blocked.json"
    blocked.write_text(json.dumps(payload), encoding="utf-8")
    code = main(
        [
            "value",
            "KO",
            "2026-10-11",
            "--assumptions",
            "examples/assumptions/coca-cola.yaml",
            "--bundle",
            str(blocked),
            "--market",
            "",
        ]
    )
    assert code == 1
    assert "Readiness blocks valuation" in capsys.readouterr().err
