"""Cross-check the DCF engine against Damodaran's free FCFF template (fcffsimpleginzu).

The template's example company and inputs are used unchanged, except its
documented reinvestment-lag override set to 0 (reinvestment funds the same
year's growth, as in this engine). LibreOffice recalculates the template; the
engine then values the same path (end-of-year discounting, the template's
cost-of-capital fade) and the results must agree within 0.5%. The inputs and
template results are saved as a fixture for offline regression tests.
"""

import argparse
import json
import tempfile
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook

from dcf_agent.engine.dcf import Drivers, Shares, dilute, project
from dcf_agent.excel.parity import recalculate

ROOT = Path(__file__).resolve().parents[1]
URL = "https://pages.stern.nyu.edu/~adamodar/pc/fcffsimpleginzu.xlsx"
TEMPLATE = ROOT / "data/raw/damodaran/templates/fcffsimpleginzu.xlsx"
FIXTURE = ROOT / "tests/fixtures/damodaran/ginzu-template.json"
TOLERANCE = 0.005
YEARS = "CDEFGHIJKL"  # explicit years 1-10 on the template's output sheet


def template_case(path: Path) -> dict[str, Any]:
    with tempfile.TemporaryDirectory() as work:
        prepared = Path(work) / "ginzu.xlsx"
        wb = load_workbook(path)
        inputs = wb["Input sheet"]
        inputs["B56"] = "Yes"  # override the default one-year reinvestment lag ...
        inputs["B57"] = 0  # ... with no lag
        wb.save(prepared)
        recalculated = recalculate(prepared, Path(work) / "out")
        values = load_workbook(recalculated, data_only=True)
    out, inp = values["Valuation output"], values["Input sheet"]

    def row(r: int) -> list[float]:
        return [float(out[f"{c}{r}"].value) for c in YEARS]

    valued = inp["B3"].value
    return {
        "source": URL,
        "company": inp["B4"].value,
        "valuation_date": valued.date().isoformat() if hasattr(valued, "date") else str(valued),
        "inputs": {
            "base_revenue": float(out["B3"].value),
            "revenue_growth": row(2),
            "operating_margin": row(4),
            "tax_rate": row(6),
            "sales_to_capital": row(38),
            "cost_of_capital": row(12),
            "terminal_growth": float(out["M2"].value),
            "terminal_margin": float(out["M4"].value),
            "terminal_tax": float(out["M6"].value),
            "terminal_roic": float(out["M40"].value),
            "terminal_wacc": float(out["B17"].value),
            "risk_free_rate": float(inp["B34"].value),
            "debt": float(out["B25"].value),
            "minority": float(out["B26"].value),
            "cash": float(out["B27"].value),
            "non_operating": float(out["B28"].value),
            "shares": float(out["B32"].value),
        },
        "template": {
            "pv_explicit": float(out["B20"].value),
            "terminal_value": float(out["B18"].value),
            "pv_terminal": float(out["B19"].value),
            "value_of_operating_assets": float(out["B24"].value),
            "equity_value": float(out["B31"].value),
            "value_per_share": float(out["B33"].value),
        },
    }


def engine_results(case: dict[str, Any]) -> dict[str, float]:
    i = case["inputs"]
    valued = date.fromisoformat(case["valuation_date"])
    if abs(i["terminal_margin"] - i["operating_margin"][-1]) > 1e-12:
        raise SystemExit("Template terminal margin differs from year 10; not an identical case")
    d = Drivers(
        valuation_date=valued,
        base_revenue=i["base_revenue"],
        # Full years from the valuation date: the template has no stub period.
        year_ends=tuple(valued + timedelta(days=365 * k) for k in range(1, 11)),
        revenue_growth=tuple(i["revenue_growth"]),
        operating_margin=tuple(i["operating_margin"]),
        sales_to_capital=tuple(i["sales_to_capital"]),
        tax_rate=tuple(i["tax_rate"]),
        terminal_growth=i["terminal_growth"],
        terminal_roic=i["terminal_roic"],
        marginal_tax_rate=i["terminal_tax"],
        wacc=i["cost_of_capital"][0],
        terminal_wacc=i["terminal_wacc"],
        risk_free_rate=i["risk_free_rate"],
        mid_year=False,
        wacc_path=tuple(i["cost_of_capital"]),
    )
    p = project(d)
    bridge = i["cash"] + i["non_operating"] - i["debt"] - i["minority"]
    equity = dilute(p.enterprise_value + bridge, Shares(basic=i["shares"]))
    return {
        "pv_explicit": p.pv_explicit,
        "terminal_value": p.terminal_value,
        "pv_terminal": p.pv_terminal,
        "value_of_operating_assets": p.enterprise_value,
        "equity_value": equity.equity_value,
        "value_per_share": equity.value_per_share,
    }


def compare(case: dict[str, Any]) -> list[tuple[str, float, float, bool]]:
    engine = engine_results(case)
    rows = []
    for key, expected in case["template"].items():
        got = engine[key]
        rows.append((key, expected, got, abs(got - expected) <= TOLERANCE * abs(expected)))
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--refresh", action="store_true", help="Download the template again")
    args = parser.parse_args()
    if args.refresh or not TEMPLATE.exists():
        response = httpx.get(URL, follow_redirects=True, timeout=60)
        response.raise_for_status()
        TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
        TEMPLATE.write_bytes(response.content)
    case = template_case(TEMPLATE)
    rows = compare(case)
    for key, expected, got, ok in rows:
        print(
            f"{'OK ' if ok else 'BAD'} {key:28} template {expected:>16,.4f} engine {got:>16,.4f} "
            f"diff {(got - expected) / expected:+.4%}"
        )
    FIXTURE.parent.mkdir(parents=True, exist_ok=True)
    FIXTURE.write_text(json.dumps(case, indent=2), encoding="utf-8")
    print(f"Saved {FIXTURE}")
    if not all(ok for *_, ok in rows):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
